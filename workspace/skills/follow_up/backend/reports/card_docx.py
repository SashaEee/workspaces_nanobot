"""
Follow Up 2.0 — экспорт карточки «Контроль исполнения» в DOCX.

Рабочая программа проверки исполнения: поручения с анализом ответа,
методология и отклонения исходной проверки, репозиторий с готовностью
к репроверке, смежные кейсы, статистика, чек-лист плана с отметками.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Dict, List, Optional

VERDICT_RU = {"solved": "решено", "not_solved": "не решено",
              "rework": "доработка"}
FORMALITY_RU = {"substantive": "содержательный", "partial": "частичный",
                "formal": "формальный (отписка)"}
READINESS_RU = {"green": "готов к репроверке",
                "amber": "частично готов — потребуется доработка",
                "red": "воспроизвести по репозиторию нельзя"}


def _fmt_rub(v) -> str:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return ""
    if v >= 1e9:
        return f"{v / 1e9:.1f} млрд ₽"
    if v >= 1e6:
        return f"{v / 1e6:.1f} млн ₽"
    if v >= 1e3:
        return f"{v / 1e3:.0f} тыс ₽"
    return f"{v:.0f} ₽"


def _method_value(v) -> Optional[str]:
    if isinstance(v, dict) and v.get("value"):
        val = v["value"]
        return ", ".join(val) if isinstance(val, list) else str(val)
    return None


def build_card_docx(card: Dict, checklist: Dict) -> bytes:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    payloads = card.get("payloads") or {}
    km = card.get("km_id") or "—"

    doc.add_heading(f"Контроль исполнения поручений — КМ {km}", level=0)
    doc.add_paragraph(
        f"Сформировано Follow Up 2.0 · {date.today().isoformat()} · "
        f"карточка {card.get('card_id') or '—'}. "
        f"Материал аналитический: вердикт принимает аудитор.")

    # ── 1. Поручения и анализ ответа ──
    poruch = payloads.get("poruch") or {}
    rows: List[Dict] = poruch.get("rows") or []
    a_by_key = {i.get("poruch_key"): i
                for i in ((poruch.get("analysis") or {}).get("items") or [])}
    doc.add_heading("1. Поручения и анализ ответа профильного подразделения",
                    level=1)
    if poruch.get("letter_reuse"):
        lr = poruch["letter_reuse"]
        doc.add_paragraph(
            f"Внимание: текст ответа почти дословно совпадает с ответом по "
            f"КМ-{lr.get('km_id')} (сходство {lr.get('similarity')}) — "
            f"признак типовой отписки.")
    for n, r in enumerate(rows, 1):
        doc.add_heading(
            f"Поручение {n} — рег.№ {r.get('doc_reg_num') or '—'} "
            f"({r.get('poruch_status') or 'нет статуса'})", level=2)
        for label, key in (("Проблема из акта", "problem"),
                           ("Текст поручения", "assignment_"),
                           ("Отчёт профильного подразделения", "actions"),
                           ("Исполнитель", "block_unit"),
                           ("Закрыто", "close_fact")):
            if r.get(key):
                p = doc.add_paragraph()
                p.add_run(f"{label}: ").bold = True
                p.add_run(str(r[key]))
        a = a_by_key.get(r.get("poruch_key"))
        if a:
            p = doc.add_paragraph()
            p.add_run("Анализ ответа: ").bold = True
            p.add_run(f"{a.get('claim') or ''} "
                      f"(характер: {FORMALITY_RU.get(a.get('formality'), '—')}, "
                      f"доказательства: {a.get('evidence_quality') or '—'})")
            if a.get("evidence"):
                doc.add_paragraph("Доказательства: " + "; ".join(a["evidence"]))
            for title, key in (("Чего не хватает", "missing_evidence"),
                               ("Запросить у профильника", "what_to_request"),
                               ("Расхождения", "discrepancies")):
                if a.get(key):
                    doc.add_paragraph(f"{title}:")
                    for x in a[key]:
                        doc.add_paragraph(str(x), style="List Bullet")
    note = (poruch.get("analysis") or {}).get("overall_note")
    if note:
        doc.add_paragraph(f"Общее наблюдение: {note}")

    # ── 2. Исходная проверка ──
    method = payloads.get("method") or {}
    doc.add_heading("2. Исходная проверка (методология и отклонения)", level=1)
    if not method.get("found"):
        doc.add_paragraph(method.get("note") or "Акт не проиндексирован.")
    else:
        if method.get("act_note"):
            doc.add_paragraph(method["act_note"])
        m = method.get("method") or {}
        for label, key in (("Периметр", "perimeter"),
                           ("Источники данных", "data_sources"),
                           ("Техники", "techniques"), ("Выборка", "sample"),
                           ("Критерий нарушения", "criteria"),
                           ("Период", "period")):
            val = _method_value(m.get(key))
            if val:
                p = doc.add_paragraph()
                p.add_run(f"{label}: ").bold = True
                p.add_run(val)
        if m.get("gaps"):
            doc.add_paragraph("Пробелы акта:")
            for g in m["gaps"]:
                doc.add_paragraph(str(g), style="List Bullet")
        devs = method.get("deviations") or []
        if devs:
            ds = method.get("dev_summary") or {}
            doc.add_heading(
                f"Отклонения исходной проверки — {ds.get('total', len(devs))} "
                f"(критичных: {ds.get('critical', '—')}"
                + (f", ущерб {_fmt_rub(ds['impact_rub'])}"
                   if ds.get("impact_rub") else "") + ")", level=2)
            for i, d in enumerate(devs, 1):
                p = doc.add_paragraph(style="List Number")
                p.add_run(f"[{d.get('severity') or 'без критичности'}] ").bold = True
                p.add_run((d.get("description") or "").strip())
                extras = []
                if d.get("poruch_ref"):
                    extras.append(f"закрывается поручением {d['poruch_ref']}")
                if d.get("financial_impact_rub"):
                    extras.append(_fmt_rub(d["financial_impact_rub"]))
                if d.get("recommendation"):
                    extras.append(f"предписание: {d['recommendation']}")
                if extras:
                    p.add_run(" (" + "; ".join(extras) + ")").italic = True

    # ── 3. Репозиторий ──
    repo = payloads.get("repo") or {}
    doc.add_heading("3. Репозиторий проверки", level=1)
    if not repo.get("found"):
        doc.add_paragraph(repo.get("note") or "Репозиторий не найден.")
    else:
        from backend.config import get_settings
        proj = get_settings().bitbucket_project or "bitbucket"
        doc.add_paragraph(
            f"bitbucket / {proj} / {repo.get('repo_slug')} "
            f"(уровень оформления {repo.get('tier') or 'C'})"
            + (f" — {repo['repo_note']}" if repo.get("repo_note") else ""))
        rd = repo.get("readiness")
        if rd:
            p = doc.add_paragraph()
            p.add_run("Готовность к репроверке: ").bold = True
            p.add_run(READINESS_RU.get(rd.get("verdict"), rd.get("verdict") or "—"))
            for nline in (rd.get("notes") or []):
                doc.add_paragraph(str(nline), style="List Bullet")
            uncovered = [c for c in (rd.get("coverage") or [])
                         if not c.get("covered")]
            if uncovered:
                doc.add_paragraph(
                    "Не покрыты скриптами поручения: "
                    + ", ".join(str(c.get("poruch_ref")) for c in uncovered))
        for f in (repo.get("files") or [])[:25]:
            line = f.get("file_path") or ""
            if f.get("descr"):
                line += f" — {f['descr']}"
            doc.add_paragraph(line, style="List Bullet")

    # ── 4. Смежные кейсы ──
    related = payloads.get("related") or {}
    items = related.get("items") or []
    doc.add_heading("4. Смежные кейсы", level=1)
    if related.get("family_note"):
        doc.add_paragraph(related["family_note"])
    if not items:
        doc.add_paragraph("Смежных кейсов не найдено.")
    for g in items:
        doc.add_paragraph(
            f"{g.get('problem_short') or ''}… (сходство {g.get('similarity')}"
            + (f", отклонений в актах: {g['dev_count']}"
               if g.get("dev_count") else "") + ")").runs[0].bold = True
        for c in (g.get("cases") or []):
            line = (f"КМ-{c.get('km_id')} — {c.get('poruch_status') or 'нет статуса'}"
                    + (f", вердикт аудитора: {VERDICT_RU.get(c.get('auditor_verdict'), c.get('auditor_verdict'))}"
                       if c.get("auditor_verdict") else ""))
            doc.add_paragraph(line, style="List Bullet")
            if c.get("actions"):
                doc.add_paragraph(f"Ответ подразделения: {c['actions']}")

    # ── 5. Историческая статистика ──
    stats = payloads.get("stats") or {}
    u = stats.get("unit_stats") or {}
    doc.add_heading("5. Историческая статистика", level=1)
    doc.add_paragraph(
        f"Подразделение: {', '.join(stats.get('units') or []) or '—'}. "
        f"Поручений: {u.get('total', '—')}, исполнено: {u.get('done', '—')}"
        + (f" ({round(u['done_share'] * 100)}% при среднем по банку "
           f"{round(u['bank_done_share'] * 100)}%)"
           if u.get("done_share") is not None
           and u.get("bank_done_share") is not None else ""))
    rp = stats.get("response_pattern")
    if rp and rp.get("formal_share") is not None:
        doc.add_paragraph(
            f"Паттерн ответов: из {rp['analyzed']} проанализированных ответов "
            f"{round(rp['formal_share'] * 100)}% — формальные или без доказательств.")
    recs = stats.get("recurrences") or []
    if recs:
        doc.add_paragraph("Рецидивы (похожие проблемы этого подразделения):")
        for r in recs:
            doc.add_paragraph(
                f"КМ-{r.get('km_id')} — {r.get('problem_short') or ''}… "
                f"({r.get('poruch_status') or 'нет статуса'}"
                + (f", закрыто {str(r['close_fact'])[:10]}"
                   if r.get("close_fact") else "") + ")",
                style="List Bullet")

    # ── 6. План проверки (чек-лист) ──
    plan = payloads.get("plan") or {}
    steps = plan.get("steps") or []
    doc.add_heading("6. План проверки исполнения", level=1)
    for i, s in enumerate(steps):
        st = checklist.get(str(i)) or {}
        mark = "[x]" if st.get("done") else "[ ]"
        p = doc.add_paragraph()
        p.add_run(f"{mark} Шаг {i + 1}. ").bold = True
        p.add_run(s.get("text") or "")
        extras = []
        if s.get("source"):
            extras.append(f"источник: {s['source']}")
        if s.get("artifact"):
            extras.append(f"файл: {s['artifact']}")
        if s.get("dev_ref"):
            extras.append(f"отклонение {s['dev_ref']}")
        if extras:
            p.add_run(" (" + "; ".join(extras) + ")").italic = True
        if st.get("note"):
            doc.add_paragraph(f"Заметка аудитора: {st['note']}")
    if plan.get("risks"):
        doc.add_heading("Риск-сигналы", level=2)
        for r in plan["risks"]:
            doc.add_paragraph(str(r), style="List Bullet")

    for style_name in ("Normal",):
        try:
            doc.styles[style_name].font.size = Pt(11)
        except Exception:
            pass

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
